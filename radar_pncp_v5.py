"""
Radar PNCP — captação diária de licitações (v5)

Melhorias da v2:
  1. Marca licitações NOVAS (coluna primeiraColeta nunca é sobrescrita) + aba "Novas" no Excel
  2. Pontuação por relevância (termos fortes/médios/fracos) e lista de exclusão (compra de bens)
  3. Guarda TODAS as palavras-chave encontradas, não só a primeira
  4. Link do PNCP montado automaticamente quando o link de origem vem vazio; valor 0 = "não informado"
  5. Retomada: se a coleta cair no meio, a próxima execução do mesmo dia continua de onde parou
  6. Coluna "dias_restantes" e ordenação por urgência
  7. Reclassificação usa o contexto do órgão e da unidade, igual à coleta
  8. Reutiliza conexão HTTP e registra tempos por estado
  9. Inclui situações suspensas mesmo quando o prazo está próximo do fim
 10. Usa informação complementar e novos serviços ambientais na classificação
 11. Lê e preserva a marcação "Já captei?" entre execuções
"""
import argparse
import json
import logging
import re
import sqlite3
import time
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
import requests

# ============================================================================
# CONFIGURAÇÃO
# ============================================================================

URL = "https://pncp.gov.br/api/consulta/v1/contratacoes/proposta"
BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "licitacoes_progaia.db"
PASTA_PLANILHAS = BASE_DIR / "planilhas"
DIAS_MINIMOS = 7

# Termos e pesos: 3 = específico do nosso nicho, 2 = relevante, 1 = genérico (precisa de contexto)
TERMOS = {
    # --- peso 3: licenciamento / estudos ambientais
    "licenciamento ambiental": 3, "estudo de impacto ambiental": 3, "estudos ambientais": 3,
    "eia/rima": 3, "eia": 3, "rima": 3, "plano básico ambiental": 3,
    "relatório de controle ambiental": 3, "plano de controle ambiental": 3,
    "autorização de supressão de vegetação": 3, "asv": 3, "evtea": 3,
    "relatório ambiental simplificado": 3, "ras": 3,
    "prad": 3, "plano de recuperação de área degradada": 3,
    "recuperação de área degradada": 3, "recuperação de áreas degradadas": 3,
    "supressão de vegetação": 3, "supressão arbórea": 3, "supressão total de árvores": 3,
    "recomposição florística": 3, "restauração florestal": 3, "inventário florestal": 3,
    "cadastro arbóreo": 3, "reposição florestal": 3,
    "manejo arbóreo": 3, "manejo de arborização": 3, "poda de árvores": 3,
    "poda": 3, "destoca": 3, "remoção de árvores": 3, "corte raso": 3,
    "supressão e poda": 3, "salvamento arqueológico": 3,
    "hidrossemeadura": 3, "revegetação": 3, "plantio de mudas": 3,
    "recomposição vegetal": 3, "espeleologia": 3, "espeleológico": 3,
    "resgate de fauna": 3,
    # Termos observados em contratos da Progaia usados para ampliar a cobertura ambiental
    "arborização urbana": 3, "manejo de árvores": 3, "avaliação de árvores": 2,
    "plantio de árvores": 3, "mudas arbóreas": 2, "coroamento de mudas": 2,
    "restauração ecológica": 3, "enleivamento": 2, "tomografia de árvores": 3,
    "tomografia sônica": 3, "inspeção subaquática": 2,
    "licença de operação": 3, "unidade de conservação": 2,
    "plano de manejo de fauna": 3, "plano de fauna": 3,
    "supervisão ambiental": 3, "supervisão ambiental de obras": 3,
    "gestão ambiental hidroviária": 3, "monitoramento hidroviário": 3,
    "levantamento hidroviário": 3, "eclusas": 3,
    "arqueologia": 3, "acompanhamento arqueológico": 3,
    # --- peso 2: gestão/monitoramento e levantamentos técnicos
    "gestão ambiental": 2, "gerenciamento ambiental": 2, "programas ambientais": 2,
    "manutenção de áreas verdes": 2, "manutenção de áreas verdes e jardinagem": 2,
    "supervisão de obras": 2, "supervisão das obras": 2, "supervisão de obra": 2,
    "assessoramento": 2, "projeto básico e executivo de engenharia": 2,
    "estudos e projetos de engenharia": 2, "projeto de engenharia": 2,
    "plano de manejo": 2, "fauna silvestre": 2,
    "monitoramento ambiental": 2, "monitoramento de fauna": 2, "monitoramento de flora": 2,
    "levantamento de fauna": 2, "impactos ambientais": 2,
    "batimetria": 2, "hidrologia": 2, "dragagem": 2,
    "aerolevantamento": 2, "rpa/vant": 2, "levantamento aerofotogramétrico": 2,
    # --- peso 1: genéricos (muito falso positivo sozinhos)
    "topografia": 1, "sondagem": 1, "projeto executivo de engenharia": 1,
    "educação ambiental": 1,
}
PALAVRAS_CURTAS = {"eia", "rima", "asv", "ras", "prad"}

# Se o objeto parecer COMPRA DE BENS, sem termo forte (peso>=3) e sem sinal de serviço/obra (ver RESGATE), é descartado.
EXCLUSOES = [
    r"\baquisi[çc][ãa]o\b", r"\bfornecimento de (?:material|materiais|uniforme|equipamento|bens)",
    r"material de consumo", r"material de expediente", r"\buniformes?\b", r"decora[çc][ãa]o",
    r"loca[çc][ãa]o de estrutura", r"g[êe]neros aliment", r"licen[çc]as? de softwares?",
    r"\bkit(?:s)?\b", r"brinquedos", r"material did[áa]tico", r"material gr[áa]fico",
]
SCORE_MINIMO = 2  # Um termo genérico isolado não basta para manter a oportunidade.
STATUS_CAPTACAO = {"nao": 0, "sim": 1, "sem interesse": 2}
ROTULOS_CAPTACAO = {0: "Não", 1: "Sim", 2: "Sem interesse"}

# Estes termos são apenas pistas de um levantamento topográfico. Não ativar sem serviço-alvo.
TOPOGRAFIA_FAMILIA = [
    "topografia", "topográfico", "topográfica", "levantamento topográfico", "agrimensura",
    "agrimensor", "planialtimétrico", "georreferenciamento", "geodésico", "aerolevantamento",
    "levantamento aerofotogramétrico", "rpa/vant",
]

# Sugestões derivadas dos serviços da Progaia; mantidas inativas até decisão do usuário.
SUGESTOES_NOVOS_TERMOS = [
    "cavidade natural subterrânea", "levantamento hidrográfico", "projeto urbanístico",
]

MODALIDADES = [4, 5, 6, 7, 8]
NOMES_MODALIDADES = {
    4: "Concorrência - Eletrônica", 5: "Concorrência - Presencial",
    6: "Pregão - Eletrônico", 7: "Pregão - Presencial", 8: "Dispensa de Licitação",
}
UFS = ["AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS", "MG", "PA", "PB",
       "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC", "SP", "SE", "TO"]

MAX_FALHAS_CONSECUTIVAS = 8
DELAY_ENTRE_REQUISICOES = 1.0

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("radar_pncp")
SESSAO_HTTP = requests.Session()

# Se o texto deixa claro que é contratação de SERVIÇO/OBRA, a exclusão não vale
RESGATE = [
    r"servi[çc]os? t[ée]cnic", r"elabora[çc][ãa]o", r"contrata[çc][ãa]o (?:semi-?)?integrada",
    r"execu[çc][ãa]o de obras?", r"obras? de engenharia", r"empresas? (?:especializada|de engenharia)",
]
_REGEX_EXCLUSAO = re.compile("|".join(EXCLUSOES), re.IGNORECASE)
_REGEX_RESGATE = re.compile("|".join(RESGATE), re.IGNORECASE)


# ============================================================================
# CLASSIFICAÇÃO
# ============================================================================

TERMOS_CONTEXTUAIS = {"supervisão de obras", "supervisão das obras", "supervisão de obra",
                      "assessoramento", "projeto básico e executivo de engenharia",
                      "estudos e projetos de engenharia", "projeto de engenharia"}
_CONTEXTO_INFRA = re.compile(r"\b(?:rodovia|br\s*-\s*\d+|ponte|ferrovia|hidrovia|duplica[çc][ãa]o|pavimenta[çc][ãa]o|rs[ct]?\s*-\s*\d+)\b", re.I)
_CONTEXTO_PLANO_MAN = re.compile(r"\b(?:ambiental|ambientais|florestal|florestais|fauna|flora|unidade de conservacao|area protegida)\b", re.I)

def _contem(texto, termo):
    texto_norm = _normalizar(texto)
    termo_norm = _normalizar(termo)
    palavras = termo_norm.split()
    tokens = [re.escape(p) + (r"(?:es)?" if p.endswith("r") else r"s?") for p in palavras]
    padrao = r"\b" + r"\s+".join(tokens) + r"\b"
    if termo in PALAVRAS_CURTAS:
        return re.search(padrao, texto_norm) is not None
    return re.search(padrao, texto_norm) is not None


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "").lower())
    return "".join(c for c in texto if not unicodedata.combining(c))


def _texto_para_classificar(objeto_compra, informacao_complementar=None):
    """Combina o objeto e o detalhe complementar, quando a API os fornece."""
    partes = []
    for valor in (objeto_compra, informacao_complementar):
        texto = str(valor or "").strip()
        if texto and texto not in partes:
            partes.append(texto)
    return " ".join(partes)


def _tem_familia_topografia(texto):
    return any(_contem(texto, termo) for termo in TOPOGRAFIA_FAMILIA)


def classificar(objeto_compra, detalhado=False, orgao=None):
    """Classifica objeto. O modo padrão mantém a tupla legada para os chamadores atuais."""
    texto = str(objeto_compra or "")
    contexto_orgao = bool(re.search(r"\b(?:dnit|der|daer|sinfra)\b|departamento nacional de infraestrutura de transportes",
                                    _normalizar(orgao or "")))
    contexto_infra = contexto_orgao or bool(_CONTEXTO_INFRA.search(_normalizar(texto)))
    achados = [t for t in TERMOS if _contem(texto, t) and
               (t not in TERMOS_CONTEXTUAIS or contexto_infra) and
               (t != "plano de manejo" or bool(_CONTEXTO_PLANO_MAN.search(_normalizar(texto))))]
    # "eia" já está dentro de "eia/rima": não conta duas vezes
    if "eia/rima" in achados:
        achados = [t for t in achados if t not in ("eia", "rima")]
    tem_topografia = _tem_familia_topografia(texto)
    # Testar cada termo individualmente para excluir toda a família topográfica do score.
    termos_score = [t for t in achados if not any(_contem(t, familia) for familia in TOPOGRAFIA_FAMILIA)]
    score = sum(TERMOS[t] for t in termos_score)
    if "programas ambientais" in termos_score and contexto_infra:
        score += 1
    tem_forte = any(TERMOS[t] >= 3 for t in termos_score) or ("programas ambientais" in termos_score and contexto_infra)
    manter = bool(termos_score) and score >= SCORE_MINIMO
    motivo = ""
    if tem_topografia and not any(TERMOS[t] >= 2 for t in termos_score):
        manter = False
        motivo = "topografia/agrimensura sem serviço-alvo"
    if not manter and not motivo:
        motivo = "sem termo de relevância"
    if not tem_forte and _REGEX_EXCLUSAO.search(texto) and not _REGEX_RESGATE.search(texto):
        manter = False
        if not motivo:
            motivo = "compra de bens sem termo forte/serviço"
    principais = termos_score or achados
    principal = max(principais, key=lambda t: (TERMOS[t], len(t))) if principais else ""
    resultado = {"manter": manter, "score": score, "termos": achados, "principal": principal, "motivo": motivo}
    if detalhado:
        return resultado
    return (score, achados, principal) if manter else None


def link_pncp(numero_controle):
    """'01613321000124-1-000132/2026' -> https://pncp.gov.br/app/editais/01613321000124/2026/132"""
    m = re.match(r"^(\d{14})-\d-(\d+)/(\d{4})$", numero_controle or "")
    if not m:
        return None
    cnpj, seq, ano = m.groups()
    return f"https://pncp.gov.br/app/editais/{cnpj}/{ano}/{int(seq)}"


# ============================================================================
# BANCO DE DADOS
# ============================================================================

COLUNAS_NOVAS = {
    "licitacoes": [("modalidadeNome", "TEXT"), ("palavrasChave", "TEXT"),
                   ("score", "INTEGER"), ("primeiraColeta", "TEXT"),
                   ("situacao", "TEXT"), ("dataAberturaProposta", "TEXT"),
                   ("dataAtualizacaoPNCP", "TEXT"), ("ultimaAlteracao", "TEXT"),
                   ("resumoAlteracoes", "TEXT"), ("descartada", "INTEGER DEFAULT 0"),
                   ("motivoDescarte", "TEXT"), ("jaCaptou", "INTEGER DEFAULT 0")],
    "log_execucoes": [("total_escaneados", "INTEGER DEFAULT 0")],
}


def iniciar_banco(db_path=None):
    conn = sqlite3.connect(db_path or DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS licitacoes (
            numeroControlePNCP TEXT PRIMARY KEY, palavraChaveEncontrada TEXT, objetoCompra TEXT,
            uf TEXT, municipio TEXT, orgao TEXT, modalidadeId INTEGER, modalidadeNome TEXT,
            dataEncerramentoProposta TEXT, valorTotalEstimado REAL, linkSistemaOrigem TEXT,
            dataColeta TEXT, dadosCompletos TEXT
        )""")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS log_execucoes (
            timestamp TEXT, total_escaneados INTEGER, total_encontradas INTEGER,
            total_falhas INTEGER, interrompido_por_circuito INTEGER
        )""")
    # checkpoint: combinações (dia, uf, modalidade) já coletadas por completo
    conn.execute("""
        CREATE TABLE IF NOT EXISTS progresso (
            dia TEXT, uf TEXT, modalidade INTEGER, PRIMARY KEY (dia, uf, modalidade)
        )""")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS historico_alteracoes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, numeroControlePNCP TEXT NOT NULL,
            data TEXT NOT NULL, campo TEXT NOT NULL, valorAntigo TEXT, valorNovo TEXT
        )""")
    for tabela, colunas in COLUNAS_NOVAS.items():
        existentes = [r[1] for r in conn.execute(f"PRAGMA table_info({tabela})")]
        for coluna, tipo in colunas:
            if coluna not in existentes:
                conn.execute(f"ALTER TABLE {tabela} ADD COLUMN {coluna} {tipo}")
    # Registros antigos: a 1ª coleta é a única data que existia
    conn.execute("UPDATE licitacoes SET primeiraColeta = dataColeta WHERE primeiraColeta IS NULL")
    conn.commit()
    return conn


def reclassificar_banco(conn):
    """Reaplica filtro sem apagar linhas e atualiza a auditoria de descartes."""
    linhas = conn.execute(
        "SELECT numeroControlePNCP, objetoCompra, orgao, dadosCompletos FROM licitacoes"
    ).fetchall()
    descartadas = 0
    for numero, objeto, orgao, dados_completos in linhas:
        try:
            dados = json.loads(dados_completos or "{}")
        except (TypeError, ValueError):
            dados = {}
        unidade = dados.get("unidadeOrgao") or {}
        contexto_orgao = " ".join(
            str(valor or "") for valor in (orgao, unidade.get("nomeUnidade"))
        )
        texto_classificacao = _texto_para_classificar(
            objeto, dados.get("informacaoComplementar")
        )
        r = classificar(texto_classificacao, detalhado=True, orgao=contexto_orgao)
        if not r["manter"]:
            conn.execute("UPDATE licitacoes SET descartada=1, motivoDescarte=? WHERE numeroControlePNCP=?",
                         (r["motivo"], numero))
            descartadas += 1
        else:
            conn.execute("""UPDATE licitacoes SET score=?, palavrasChave=?, palavraChaveEncontrada=?,
                         descartada=0, motivoDescarte=NULL WHERE numeroControlePNCP=?""",
                         (r["score"], "; ".join(r["termos"]), r["principal"], numero))
    conn.commit()
    log.info(f"Reclassificação: {len(linhas) - descartadas} mantidas | {descartadas} marcadas como descartadas")
    return descartadas


def salvar_licitacao(conn, item, classificacao):
    numero = item.get("numeroControlePNCP")
    if not numero:
        return None
    score, termos, principal = classificacao
    try:
        valor = float(item["valorTotalEstimado"]) if item.get("valorTotalEstimado") is not None else None
    except (ValueError, TypeError):
        valor = None
    modalidade_id = item.get("modalidadeId")
    hoje = date.today().isoformat()
    unidade = item.get("unidadeOrgao") or {}
    orgao = item.get("orgaoEntidade") or {}
    situacao = item.get("situacaoCompraNome")
    data_abertura = item.get("dataAberturaProposta")
    data_atualizacao = item.get("dataAtualizacao")
    novo = {
        "dataAberturaProposta": data_abertura,
        "dataEncerramentoProposta": item.get("dataEncerramentoProposta"),
        "situacao": situacao,
        "valorTotalEstimado": valor,
        "linkSistemaOrigem": item.get("linkSistemaOrigem"),
        "objetoCompra": item.get("objetoCompra"),
        "dataAtualizacaoPNCP": data_atualizacao,
    }
    atual = conn.execute("SELECT * FROM licitacoes WHERE numeroControlePNCP=?", (numero,)).fetchone()
    ja_existia = atual is not None
    alteracoes = []
    if atual:
        colunas = [r[1] for r in conn.execute("PRAGMA table_info(licitacoes)")]
        atual = dict(zip(colunas, atual))
        for campo, valor_novo in novo.items():
            # dataAtualizacao do PNCP pode avançar sem mudança material; persistimos o campo,
            # mas não registramos isso como alteração da licitação.
            if campo == "dataAtualizacaoPNCP":
                continue
            valor_antigo = atual.get(campo)
            if campo == "valorTotalEstimado":
                # A API usa 0 como valor ausente: não sobrescrever valor conhecido com zero.
                if valor_novo in (None, 0):
                    valor_novo = valor_antigo
            elif valor_novo is None and valor_antigo is not None:
                # Campos ausentes da resposta não apagam a última informação conhecida.
                valor_novo = valor_antigo
            if valor_novo != valor_antigo:
                alteracoes.append((campo, valor_antigo, valor_novo))
        for campo, antigo, novo_valor in alteracoes:
            conn.execute("INSERT INTO historico_alteracoes (numeroControlePNCP,data,campo,valorAntigo,valorNovo) VALUES (?,?,?,?,?)",
                         (numero, hoje, campo, _valor_texto(antigo), _valor_texto(novo_valor)))
        resumo = atual.get("resumoAlteracoes") or ""
        if alteracoes:
            trechos = [f"{campo}: {_resumir_valor(a)} → {_resumir_valor(b)}" for campo, a, b in alteracoes]
            linha = f"{datetime.now().strftime('%d/%m/%Y')}: " + "; ".join(trechos)
            resumo = (resumo + " | " + linha).strip(" |")
    else:
        resumo = None
    final = dict(novo)
    if atual:
        # Persistir o valor conhecido quando a API envia zero/ausente.
        for campo, valor_novo in final.items():
            if valor_novo is None and atual.get(campo) is not None:
                final[campo] = atual[campo]
        if final["valorTotalEstimado"] in (None, 0) and atual.get("valorTotalEstimado") not in (None, 0):
            final["valorTotalEstimado"] = atual["valorTotalEstimado"]
    conn.execute("""
        INSERT INTO licitacoes (
            numeroControlePNCP, palavraChaveEncontrada, palavrasChave, score, objetoCompra, uf, municipio,
            orgao, modalidadeId, modalidadeNome, dataEncerramentoProposta, valorTotalEstimado,
            linkSistemaOrigem, dataColeta, primeiraColeta, dadosCompletos, situacao,
            dataAberturaProposta, dataAtualizacaoPNCP, ultimaAlteracao, resumoAlteracoes
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(numeroControlePNCP) DO UPDATE SET
            palavraChaveEncontrada=excluded.palavraChaveEncontrada,
            palavrasChave=excluded.palavrasChave, score=excluded.score,
            objetoCompra=excluded.objetoCompra, uf=excluded.uf, municipio=excluded.municipio,
            orgao=excluded.orgao, modalidadeId=excluded.modalidadeId, modalidadeNome=excluded.modalidadeNome,
            dataEncerramentoProposta=excluded.dataEncerramentoProposta,
            valorTotalEstimado=excluded.valorTotalEstimado, linkSistemaOrigem=excluded.linkSistemaOrigem,
            dataColeta=excluded.dataColeta, dadosCompletos=excluded.dadosCompletos,
            descartada=0, motivoDescarte=NULL,
            situacao=excluded.situacao, dataAberturaProposta=excluded.dataAberturaProposta,
            dataAtualizacaoPNCP=excluded.dataAtualizacaoPNCP,
            ultimaAlteracao=CASE WHEN excluded.ultimaAlteracao IS NOT NULL THEN excluded.ultimaAlteracao ELSE licitacoes.ultimaAlteracao END,
            resumoAlteracoes=COALESCE(excluded.resumoAlteracoes,licitacoes.resumoAlteracoes)
            -- primeiraColeta NÃO é atualizada: é ela que diz se a licitação é nova
    """, (
        numero, principal, "; ".join(termos), score, item.get("objetoCompra"),
        unidade.get("ufSigla"), unidade.get("municipioNome"),
        orgao.get("razaoSocial"), modalidade_id,
        NOMES_MODALIDADES.get(modalidade_id, f"Modalidade {modalidade_id}"),
        final["dataEncerramentoProposta"], final["valorTotalEstimado"], final["linkSistemaOrigem"],
        hoje, atual.get("primeiraColeta") if atual else hoje, json.dumps(item, ensure_ascii=False),
        final["situacao"], final["dataAberturaProposta"], final["dataAtualizacaoPNCP"],
        hoje if alteracoes else None, resumo,
    ))
    return "ATUALIZADA" if alteracoes else ("SEM ALTERACAO" if ja_existia else "NOVA")


def _valor_texto(valor):
    return None if valor is None else str(valor)


def _resumir_valor(valor):
    if valor in (None, ""):
        return "(vazio)"
    return str(valor)[:80]


def registrar_execucao(conn, escaneados, encontradas, falhas, interrompido):
    conn.execute(
        """INSERT INTO log_execucoes
           (timestamp,total_escaneados,total_encontradas,total_falhas,interrompido_por_circuito)
           VALUES (?,?,?,?,?)""",
        (pd.Timestamp.now().isoformat(timespec="seconds"), escaneados, encontradas, falhas, int(interrompido)))


# ============================================================================
# API
# ============================================================================

def consultar_pagina(pagina, codigo_modalidade, uf, data_inicial, data_final, tentativas=4, timeout=20,
                     metricas=None):
    params = {"dataInicial": data_inicial, "dataFinal": data_final, "pagina": pagina,
              "codigoModalidadeContratacao": codigo_modalidade, "uf": uf, "tamanhoPagina": 50}
    for tentativa in range(tentativas):
        try:
            inicio_requisicao = time.monotonic()
            try:
                resp = SESSAO_HTTP.get(URL, params=params, timeout=timeout)
            finally:
                if metricas is not None:
                    metricas["chamadas"] += 1
                    metricas["resposta_s"] += time.monotonic() - inicio_requisicao
            if resp.status_code == 429:
                espera = int(resp.headers.get("Retry-After", 5 * (tentativa + 1)))
                log.warning(f"429 (rate limit) em {uf}/modalidade {codigo_modalidade}. Esperando {espera}s...")
                if metricas is not None:
                    metricas["espera_429_s"] += espera
                time.sleep(espera)
                continue
            if resp.status_code == 204:
                return None
            resp.raise_for_status()
            return resp.json()
        except requests.exceptions.Timeout:
            time.sleep(2 ** tentativa)
        except requests.exceptions.RequestException as e:
            log.warning(f"Erro de requisição ({uf}, modalidade {codigo_modalidade}, pág {pagina}): {e}")
            time.sleep(2 ** tentativa)
    return "FALHOU"


# ============================================================================
# COLETA (com retomada)
# ============================================================================

def coletar(dias_futuros=90, delay=None, nova_varredura=False):
    delay = DELAY_ENTRE_REQUISICOES if delay is None else delay
    hoje = date.today()
    dia = hoje.isoformat()
    data_final = hoje + timedelta(days=dias_futuros)
    conn = iniciar_banco()
    reclassificar_banco(conn)
    if nova_varredura:
        conn.execute("DELETE FROM progresso WHERE dia = ?", (dia,))
        conn.commit()
        feitos = set()
        log.info("Nova varredura completa solicitada; os checkpoints de hoje foram reiniciados.")
    else:
        feitos = {(u, m) for u, m in conn.execute("SELECT uf, modalidade FROM progresso WHERE dia = ?", (dia,))}
    total_escaneados = total_encontradas = total_falhas = falhas_consecutivas = 0
    novas = atualizadas = descartadas = 0
    interrompido = False

    log.info(f"Buscando propostas abertas de {hoje:%d/%m/%Y} até {data_final:%d/%m/%Y}.")
    if feitos:
        log.info(f"Retomando: {len(feitos)} de {len(UFS) * len(MODALIDADES)} combinações já coletadas hoje.")

    for uf in UFS:
        if interrompido:
            break
        metricas_uf = {"chamadas": 0, "resposta_s": 0.0, "espera_429_s": 0.0, "pausa_s": 0.0}
        inicio_uf = time.monotonic()
        for modalidade in MODALIDADES:
            if (uf, modalidade) in feitos:
                continue
            pagina, combo_ok = 1, True
            while True:
                resultado = consultar_pagina(pagina, modalidade, uf, hoje.strftime("%Y%m%d"),
                                             data_final.strftime("%Y%m%d"), metricas=metricas_uf)
                metricas_uf["pausa_s"] += delay
                time.sleep(delay)
                if resultado == "FALHOU":
                    total_falhas += 1
                    falhas_consecutivas += 1
                    combo_ok = False
                    if falhas_consecutivas >= MAX_FALHAS_CONSECUTIVAS:
                        log.error(f"{MAX_FALHAS_CONSECUTIVAS} falhas consecutivas — interrompendo. "
                                  "Rode de novo: a coleta continua de onde parou.")
                        interrompido = True
                    break
                falhas_consecutivas = 0
                if resultado is None or not resultado.get("data"):
                    break
                total_escaneados += len(resultado["data"])
                for item in resultado["data"]:
                    unidade = item.get("unidadeOrgao") or {}
                    entidade = item.get("orgaoEntidade") or {}
                    texto_classificacao = _texto_para_classificar(
                        item.get("objetoCompra"), item.get("informacaoComplementar")
                    )
                    detalhes = classificar(texto_classificacao, detalhado=True,
                        orgao=" ".join(str(v or "") for v in (entidade.get("razaoSocial"), unidade.get("nomeUnidade"))))
                    if detalhes["manter"]:
                        movimento = salvar_licitacao(conn, item, (detalhes["score"], detalhes["termos"], detalhes["principal"]))
                        novas += movimento == "NOVA"
                        atualizadas += movimento == "ATUALIZADA"
                        total_encontradas += 1
                    elif detalhes["motivo"] == "topografia/agrimensura sem serviço-alvo":
                        _salvar_descartada(conn, item, detalhes["motivo"])
                        descartadas += 1
                if pagina >= resultado.get("totalPaginas", 1):
                    break
                pagina += 1
            if combo_ok:  # só marca como feita se não houve falha
                conn.execute("INSERT OR IGNORE INTO progresso VALUES (?,?,?)", (dia, uf, modalidade))
            if interrompido:
                break
        conn.commit()
        duracao_uf = time.monotonic() - inicio_uf
        log.info(f"{uf} concluído. Escaneadas: {total_escaneados} | duração: {duracao_uf:.1f}s | "
                 f"chamadas: {metricas_uf['chamadas']} ({metricas_uf['resposta_s']:.1f}s de resposta) | "
                 f"pausa normal: {metricas_uf['pausa_s']:.1f}s | espera 429: {metricas_uf['espera_429_s']:.1f}s")

    registrar_execucao(conn, total_escaneados, total_encontradas, total_falhas, interrompido)
    conn.execute("DELETE FROM progresso WHERE dia < ?", (dia,))  # limpa checkpoints antigos
    conn.commit()
    conn.close()
    log.info(f"Coleta finalizada. {total_escaneados} escaneadas | {total_encontradas} relevantes | "
             f"{total_falhas} falhas | {'INCOMPLETA' if interrompido else 'completa'}")
    log.info(f"Movimentos na coleta: {novas} novas | {atualizadas} atualizadas | "
             f"{descartadas} descartadas por topografia isolada")


def _salvar_descartada(conn, item, motivo):
    """Guarda no banco apenas a classe de falso positivo solicitada para auditoria."""
    numero = item.get("numeroControlePNCP")
    if not numero:
        return
    unidade = item.get("unidadeOrgao") or {}
    orgao = item.get("orgaoEntidade") or {}
    modalidade_id = item.get("modalidadeId")
    hoje = date.today().isoformat()
    try:
        valor = float(item["valorTotalEstimado"]) if item.get("valorTotalEstimado") is not None else None
    except (ValueError, TypeError):
        valor = None
    conn.execute("""INSERT INTO licitacoes (
        numeroControlePNCP,objetoCompra,uf,municipio,orgao,modalidadeId,modalidadeNome,
        dataEncerramentoProposta,valorTotalEstimado,linkSistemaOrigem,dataColeta,primeiraColeta,
        dadosCompletos,situacao,dataAberturaProposta,dataAtualizacaoPNCP,descartada,motivoDescarte
    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,?)
    ON CONFLICT(numeroControlePNCP) DO UPDATE SET dataColeta=excluded.dataColeta,
        dadosCompletos=excluded.dadosCompletos, descartada=1, motivoDescarte=excluded.motivoDescarte""",
        (numero,item.get("objetoCompra"),unidade.get("ufSigla"),unidade.get("municipioNome"),
         orgao.get("razaoSocial"),modalidade_id,NOMES_MODALIDADES.get(modalidade_id,f"Modalidade {modalidade_id}"),
         item.get("dataEncerramentoProposta"),valor,item.get("linkSistemaOrigem"),hoje,hoje,
         json.dumps(item,ensure_ascii=False),item.get("situacaoCompraNome"),item.get("dataAberturaProposta"),
         item.get("dataAtualizacao"),motivo))


# ============================================================================
# EXPORT PARA EXCEL
# ============================================================================

def limpar_para_excel(texto):
    if not isinstance(texto, str):
        return texto
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", texto)


def _timestamp_brasilia(valor):
    if not valor:
        return pd.NaT
    try:
        ts = pd.Timestamp(valor)
        if ts.tzinfo is not None:
            ts = ts.tz_convert("America/Sao_Paulo").tz_localize(None)
        return ts
    except (ValueError, TypeError):
        return pd.NaT


def _eh_dnit(orgao, dados_completos):
    texto = _normalizar(orgao)
    if "departamento nacional de infraestrutura de transportes" in texto or re.search(r"\bdnit\b", texto):
        return True
    try:
        dados = json.loads(dados_completos or "{}")
    except (TypeError, ValueError):
        dados = {}
    unidade = dados.get("unidadeOrgao") or {}
    nome = _normalizar(unidade.get("nomeUnidade", ""))
    return "departamento nacional de infraestrutura de transportes" in nome or bool(re.search(r"\bdnit\b", nome))


def _situacao_cancelada(valor):
    return _normalizar(valor).strip() in {"anulada", "anulado", "revogada", "revogado", "cancelada", "cancelado"}


def _situacao_suspensa(valor):
    return _normalizar(valor).strip() in {"suspensa", "suspenso"}


def _informacoes_apresentacao(dados_completos):
    """Extrai rótulos amigáveis para a planilha a partir do JSON PNCP salvo."""
    try:
        dados = json.loads(dados_completos or "{}")
    except (TypeError, ValueError):
        dados = {}
    tipo = dados.get("tipoInstrumentoConvocatorioNome") or ""
    fontes = dados.get("fontesOrcamentarias") or []
    descricoes = []
    for fonte in fontes:
        if not isinstance(fonte, dict):
            continue
        nome = str(fonte.get("nome") or "").strip()
        descricao = str(fonte.get("descricao") or "").strip()
        trecho = " — ".join(parte for parte in (nome, descricao) if parte)
        if trecho and trecho not in descricoes:
            descricoes.append(trecho)
    return tipo, " | ".join(descricoes)


def montar_dataframe(conn, hoje=None, incluir_sem_interesse=False):
    agora_brt = pd.Timestamp.now(tz="America/Sao_Paulo").tz_localize(None)
    hoje = pd.Timestamp(hoje or agora_brt.date()).normalize()
    # rowid fornece um ID curto, único e estável enquanto o registro permanecer no banco.
    df = pd.read_sql_query("SELECT rowid AS idRadar, * FROM licitacoes", conn)
    df["encerramento"] = pd.to_datetime(
        df["dataEncerramentoProposta"].map(_timestamp_brasilia), errors="coerce")
    df["descartada"] = df["descartada"].fillna(0).astype(int)
    df["dnit"] = df.apply(lambda r: "DNIT" if _eh_dnit(r.get("orgao"), r.get("dadosCompletos")) else "", axis=1)
    ainda_aberta = df["encerramento"] >= agora_brt
    df_aberto = df[(df["descartada"] == 0) & ainda_aberta].copy()
    if not incluir_sem_interesse:
        df_aberto = df_aberto[df_aberto["jaCaptou"].fillna(0).astype(int) != 2].copy()
    df_aberto["jaCaptou"] = df_aberto["jaCaptou"].fillna(0).astype(int).map(
        lambda valor: ROTULOS_CAPTACAO.get(valor, "Não"))
    df_aberto["idVisivel"] = df_aberto["idRadar"].map(
        lambda valor: f"R-{int(valor):06d}" if pd.notna(valor) else "")
    extras = df_aberto["dadosCompletos"].map(_informacoes_apresentacao)
    df_aberto["tipoPublicacao"] = extras.map(lambda valor: valor[0])
    df_aberto["fonteOrcamentaria"] = extras.map(lambda valor: valor[1])
    df_aberto["dias_restantes"] = (df_aberto["encerramento"].dt.normalize() - hoje).dt.days
    excecao_status = df_aberto["situacao"].map(
        lambda valor: _situacao_cancelada(valor) or _situacao_suspensa(valor)
    )
    df_aberto = df_aberto[(df_aberto["dias_restantes"] >= DIAS_MINIMOS) | excecao_status].copy()
    df_aberto["movimento"] = ""
    df_aberto.loc[df_aberto["primeiraColeta"] == hoje.date().isoformat(), "movimento"] = "NOVA"
    df_aberto.loc[(df_aberto["movimento"] == "") & (df_aberto["ultimaAlteracao"] == hoje.date().isoformat()), "movimento"] = "ATUALIZADA"
    ultima_completa = conn.execute("""SELECT timestamp FROM log_execucoes
        WHERE COALESCE(interrompido_por_circuito,0)=0 ORDER BY timestamp DESC LIMIT 1""").fetchone()
    if ultima_completa:
        data_execucao = str(ultima_completa[0])[:10]
        nao_vista = (df_aberto["dataColeta"].fillna("") < data_execucao) & (df_aberto["movimento"] == "")
        df_aberto.loc[nao_vista, "movimento"] = "NÃO VISTA NA ÚLTIMA COLETA"
    df_aberto["valorTotalEstimado"] = df_aberto["valorTotalEstimado"].astype(object)
    df_aberto.loc[df_aberto["valorTotalEstimado"].isna() | (df_aberto["valorTotalEstimado"] == 0), "valorTotalEstimado"] = "não informado"
    df_aberto["linkPNCP"] = df_aberto["numeroControlePNCP"].apply(link_pncp)
    df_aberto["urlDestino"] = df_aberto["linkPNCP"]
    df_aberto["abrirEdital"] = df_aberto["urlDestino"].map(
        lambda valor: "Abrir edital" if isinstance(valor, str) and valor.startswith(("http://", "https://")) else "")
    df_aberto = df_aberto.sort_values(["dnit", "dias_restantes", "score"], ascending=[False, True, False])
    colunas = ["dnit", "idVisivel", "jaCaptou", "movimento", "dias_restantes", "dataEncerramentoProposta",
               "objetoCompra", "orgao", "uf", "municipio", "tipoPublicacao", "valorTotalEstimado",
               "modalidadeNome", "fonteOrcamentaria", "score", "situacao", "dataAberturaProposta",
               "palavrasChave", "primeiraColeta", "dataColeta", "abrirEdital", "urlDestino", "numeroControlePNCP"]
    df_principal = df_aberto[colunas].copy()
    df_principal = df_principal.rename(columns={"jaCaptou": "Já captei?"})
    data_cols = ["dataEncerramentoProposta", "dataAberturaProposta", "primeiraColeta", "dataColeta"]
    for quadro in (df_principal,):
        for c in data_cols:
            quadro[c] = quadro[c].map(_timestamp_brasilia)
        for c in quadro.columns:
            if c not in data_cols:
                quadro[c] = quadro[c].map(limpar_para_excel)
    return df_principal


def importar_marcacoes_captacao(conn, caminho):
    """Importa os três estados das abas operacionais e de sem interesse."""
    from openpyxl import load_workbook

    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        importadas = 0
        for nome_aba in ("Em Aberto", "Sem interesse"):
            if nome_aba not in livro.sheetnames:
                continue
            aba = livro[nome_aba]
            cabecalhos = {str(celula.value or "").strip(): i for i, celula in enumerate(aba[1])}
            idx_controle = cabecalhos.get("numeroControlePNCP")
            idx_id_radar = cabecalhos.get("ID")
            idx_status = cabecalhos.get("Já captei?")
            if (idx_controle is None and idx_id_radar is None) or idx_status is None:
                continue
            for linha in aba.iter_rows(min_row=2, values_only=True):
                status = _normalizar(str(linha[idx_status] or "")).strip()
                if status not in STATUS_CAPTACAO:
                    continue
                valor_status = STATUS_CAPTACAO[status]
                if idx_controle is not None:
                    numero = str(linha[idx_controle] or "").strip()
                    if not numero:
                        continue
                    cur = conn.execute(
                        "UPDATE licitacoes SET jaCaptou=? WHERE numeroControlePNCP=?",
                        (valor_status, numero),
                    )
                else:
                    identificador = str(linha[idx_id_radar] or "").strip()
                    correspondencia = re.fullmatch(r"R-(\d+)", identificador)
                    if not correspondencia:
                        continue
                    cur = conn.execute(
                        "UPDATE licitacoes SET jaCaptou=? WHERE rowid=?",
                        (valor_status, int(correspondencia.group(1))),
                    )
                importadas += cur.rowcount
        conn.commit()
        return importadas
    finally:
        livro.close()


def exportar_excel(db_path=None, excel_path=None, hoje=None):
    conn = iniciar_banco(db_path)
    data_hoje = pd.Timestamp(hoje or date.today()).date()
    destino = Path(excel_path) if excel_path else Path(PASTA_PLANILHAS) / f"{data_hoje:%d.%m.%Y}.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)
    planilhas_anteriores = list(Path(PASTA_PLANILHAS).glob("*.xlsx"))
    # Ao rodar mais de uma vez no mesmo dia, a planilha de hoje é a fonte
    # das marcações, mesmo que outra planilha tenha sido modificada depois.
    planilha_origem = destino if destino.exists() else (
        max(planilhas_anteriores, key=lambda p: p.stat().st_mtime) if planilhas_anteriores else None
    )
    try:
        if planilha_origem:
            quantidade = importar_marcacoes_captacao(conn, planilha_origem)
            if quantidade:
                log.info(f"Marcações 'Já captei?' importadas de {planilha_origem.name}: {quantidade}")
        conn.commit()
        df = montar_dataframe(conn, hoje=hoje)
        df_sem_interesse = montar_dataframe(conn, hoje=hoje, incluir_sem_interesse=True)
        df_sem_interesse = df_sem_interesse[df_sem_interesse["Já captei?"] == "Sem interesse"]
        df_ja_captei = df[df["Já captei?"] == "Sim"]
    finally:
        conn.close()
    novas = df[df["movimento"] == "NOVA"]
    atualizadas = df[df["movimento"] == "ATUALIZADA"]
    nao_vistas = df[df["movimento"] == "NÃO VISTA NA ÚLTIMA COLETA"]
    dnit = df[df["dnit"] == "DNIT"]
    tabelas = [("Em Aberto", df), ("DNIT", dnit), ("Novas", novas), ("Atualizadas", atualizadas),
               ("Já captei", df_ja_captei), ("Sem interesse", df_sem_interesse)]
    try:
        _gravar_planilha(destino, tabelas)
    except PermissionError:
        raise PermissionError(f"Não foi possível salvar {destino}. Feche o arquivo no Excel e tente novamente.")
    log.info(f"Excel salvo em {destino}: {len(novas)} novas | {len(atualizadas)} atualizadas | "
             f"{len(nao_vistas)} não vistas")


def _gravar_planilha(destino, tabelas):
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    cores = {"NOVA": "C6EFCE", "ATUALIZADA": "FFEB9C", "NÃO VISTA NA ÚLTIMA COLETA": "D9D9D9"}
    cor_ja_captei = "DDEBF7"
    rotulos = {
        "dnit": "Prioridade", "idVisivel": "ID", "movimento": "Atualização",
        "dias_restantes": "Dias restantes", "dataEncerramentoProposta": "Fim das propostas",
        "objetoCompra": "Objeto", "orgao": "Órgão", "uf": "UF", "municipio": "Município",
        "tipoPublicacao": "Tipo de publicação", "valorTotalEstimado": "Valor estimado",
        "modalidadeNome": "Modalidade", "fonteOrcamentaria": "Fonte orçamentária",
        "score": "Score", "situacao": "Situação", "dataAberturaProposta": "Início das propostas",
        "palavrasChave": "Termos encontrados", "primeiraColeta": "Primeira identificação",
        "dataColeta": "Última coleta", "abrirEdital": "Abrir edital",
    }
    with pd.ExcelWriter(destino, engine="openpyxl") as writer:
        for aba_nome, quadro in tabelas:
            colunas_internas = ["numeroControlePNCP", "urlDestino"]
            quadro.drop(columns=colunas_internas, errors="ignore").rename(
                columns=rotulos).to_excel(writer, sheet_name=aba_nome, index=False)
            aba = writer.sheets[aba_nome]
            aba.freeze_panes = "A2"
            aba.auto_filter.ref = aba.dimensions
            indices = {c.value: c.column for c in aba[1]}
            if "Já captei?" in indices and aba.max_row > 1:
                coluna_captacao = indices["Já captei?"]
                validacao = DataValidation(type="list", formula1='"Não,Sim,Sem interesse"', allow_blank=False)
                validacao.error = "Escolha uma das três opções da lista."
                validacao.errorTitle = "Valor inválido"
                validacao.prompt = "Escolha Sim, Não ou Sem interesse para este edital."
                validacao.promptTitle = "Controle de captação"
                aba.add_data_validation(validacao)
                validacao.add(
                    f"{aba.cell(2, coluna_captacao).coordinate}:"
                    f"{aba.cell(aba.max_row, coluna_captacao).coordinate}"
                )
            larguras = {"Objeto": 70, "Órgão": 35, "Modalidade": 26, "Termos encontrados": 32,
                        "Prioridade": 12, "Atualização": 25, "Já captei?": 14, "ID": 12,
                        "Tipo de publicação": 22, "Fonte orçamentária": 48, "Abrir edital": 18,
                        "Dias restantes": 16, "Valor estimado": 18}
            for nome, largura in larguras.items():
                if nome in indices:
                    aba.column_dimensions[aba.cell(1, indices[nome]).column_letter].width = largura
            for linha in range(2, aba.max_row + 1):
                aba.row_dimensions[linha].height = 45
                for nome in ("Fim das propostas", "Início das propostas", "Primeira identificação", "Última coleta"):
                    if nome in indices:
                        celula_data = aba.cell(linha, indices[nome])
                        valor_data = celula_data.value
                        tem_hora = isinstance(valor_data, (datetime, pd.Timestamp)) and bool(
                            valor_data.hour or valor_data.minute or valor_data.second)
                        celula_data.number_format = "dd/mm/yyyy hh:mm" if tem_hora else "dd/mm/yyyy"
                for nome in ("Objeto", "Fonte orçamentária"):
                    if nome in indices:
                        aba.cell(linha, indices[nome]).alignment = Alignment(wrap_text=True, vertical="center")
                if "Valor estimado" in indices:
                    valor = aba.cell(linha, indices["Valor estimado"])
                    if isinstance(valor.value, (int, float)):
                        valor.number_format = 'R$ #,##0.00'
                if ("Abrir edital" in indices and "urlDestino" in quadro.columns
                        and linha - 2 < len(quadro)):
                    celula_link = aba.cell(linha, indices["Abrir edital"])
                    alvo = quadro["urlDestino"].iloc[linha - 2]
                    if isinstance(alvo, str) and alvo.startswith(("http://", "https://")):
                        celula_link.value = "Abrir edital"
                        celula_link.hyperlink = alvo
                        celula_link.style = "Hyperlink"
            if "movimento" in quadro.columns:
                idx_movimento = indices["Atualização"]
                idx_situacao = indices["Situação"]
                for row in range(2, aba.max_row + 1):
                    mov = aba.cell(row, idx_movimento).value
                    situacao = aba.cell(row, idx_situacao).value
                    cancelada = _situacao_cancelada(situacao)
                    suspensa = _situacao_suspensa(situacao)
                    destaque_vermelho = cancelada or suspensa
                    ja_captei = ("Já captei?" in indices
                                 and aba.cell(row, indices["Já captei?"]).value == "Sim")
                    cor = "FFC7CE" if destaque_vermelho else (
                        cor_ja_captei if ja_captei else cores.get(mov)
                    )
                    if cor:
                        preenchimento = PatternFill(fill_type="solid", fgColor=cor)
                        for cell in aba[row]:
                            cell.fill = preenchimento
                            if destaque_vermelho:
                                cell.font = Font(color="9C0006")
                    if "Prioridade" in indices and aba.cell(row, indices["Prioridade"]).value == "DNIT":
                        cell = aba.cell(row, indices["Prioridade"])
                        cell.font = Font(bold=True, color="9C0006" if destaque_vermelho else "FFFFFF")
                        if not destaque_vermelho:
                            cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
            grade = Side(style="thin", color="A6A6A6")
            borda = Border(left=grade, right=grade, top=grade, bottom=grade)
            for linha in aba.iter_rows():
                for celula in linha:
                    celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    celula.border = borda


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Coleta oportunidades do PNCP e exporta a planilha diária.")
    parser.add_argument(
        "--nova-varredura", action="store_true",
        help="consulta novamente todos os estados e modalidades, mesmo os já concluídos hoje",
    )
    argumentos = parser.parse_args()
    coletar(dias_futuros=90, nova_varredura=argumentos.nova_varredura)
    exportar_excel()

